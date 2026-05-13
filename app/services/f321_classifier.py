"""F-321 — classifier for Chadi's .docx tutoring archive.

Two layers:

  1. Pre-classification skip filters (no API call):
     - Filename pattern match against known publisher series
     - First-2KB content scan for publisher attributions / ISBN
     - Duplicate-suffix collapse so Phase B doesn't re-classify the
       same content as ``(1)`` ``(2)`` ``_-_Copie`` variants

  2. Haiku classifier (one API call per surviving file):
     - Filename + first ~3K chars of text → structured JSON verdict
     - Output fields: verdict, source_type, content_type, has_tables,
       has_paragraph_glosses, suggested_topic_slug, confidence, reasoning

Output of Phase B (scripts/classify_f321_docx.py) is a CSV with one
row per non-skipped file containing all classification fields. Phase C
dispatches each file to the right extractor based on these fields.

Hard exclusions (per Chadi 2026-05-13 audit lock):
  - Filename matches publisher patterns → source_type =
    third_party_publisher_DO_NOT_EXTRACT, classifier never called.
  - Content contains publisher attributions → same.

Topic slug enum (Phase 1 lock, 3 topics):
  - faux_amis | calques_anglais | prepositions_a_de_dans_par_pour |
    other (Phase 1 out-of-scope; classifier may suggest these but the
    extractor pipeline ignores chunks not in the locked trio).
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from docx import Document


# ── pre-classification skip filters ──────────────────────────────


# Filename patterns that indicate published material from
# copyrighted FR-as-second-language publishers (CLE International,
# Hachette FLE, Didier, Cambridge, etc.). Source: F-321.audit
# verdict 2026-05-13 — Vocabulaire_v2 = Claire Miquel / CLE.
_PUBLISHER_FILENAME_PATTERNS = [
    re.compile(r"progressif", re.IGNORECASE),
    re.compile(r"communication_progressive", re.IGNORECASE),
    re.compile(r"\bniveau_", re.IGNORECASE),
    re.compile(r"vocabulaire_v\d", re.IGNORECASE),
]


# Substrings in first-2KB that flag publisher attribution. These cover
# the publisher imprints + ISBN format + common Miquel co-authorship.
_PUBLISHER_CONTENT_MARKERS = [
    "CLE International",
    "Hachette FLE",
    "Hachette Livre",
    "Didier",
    "Cambridge University Press",
    "Claire Miquel",
    "Vocabulaire Progressif du Français",
    "Communication Progressive du Français",
]


# ISBN-13 pattern (978/979 prefix + 10 digits, hyphens optional).
_ISBN_PATTERN = re.compile(r"\b97[89][\d\- ]{10,17}\b")


# Duplicate-suffix patterns that indicate a copied file rather than
# distinct content. Match canonical-stem + suffix; collapse to first
# match by sorted filename.
_DUPLICATE_SUFFIX_PATTERN = re.compile(
    r"^(?P<stem>.+?)"
    r"(?:"
    r"\s*\(\d+\)|"          # foo (1).docx
    r"_-_Copie(?:\(\d+\))?|"  # foo_-_Copie.docx, foo_-_Copie(1).docx
    r"\s+-\s+Copy(?:\s+\(\d+\))?|"  # foo - Copy.docx, foo - Copy (2).docx
    r"\(copie\)|"           # foo(copie).docx
    r"_copy\d*"             # foo_copy.docx, foo_copy2.docx
    r")"
    r"\.docx$",
    re.IGNORECASE,
)


def is_publisher_import_by_filename(filename: str) -> bool:
    """True if `filename` matches a known publisher-series pattern.
    Caller short-circuits the classifier and tags
    source_type='third_party_publisher_DO_NOT_EXTRACT'."""
    return any(p.search(filename) for p in _PUBLISHER_FILENAME_PATTERNS)


def is_publisher_import_by_content(text: str) -> Optional[str]:
    """If text contains a known publisher attribution OR an ISBN
    pattern, returns the matched marker string. None otherwise.

    Caller passes the first 2KB of `text`; the function is content-
    agnostic so callers can also pass full text if needed.
    """
    for marker in _PUBLISHER_CONTENT_MARKERS:
        if marker in text:
            return marker
    isbn_match = _ISBN_PATTERN.search(text)
    if isbn_match:
        return f"ISBN:{isbn_match.group(0).strip()}"
    return None


def detect_duplicate_suffix(filename: str) -> Optional[str]:
    """If `filename` looks like a duplicate-suffix copy (`(1)`,
    `_-_Copie`, ` - Copy`, `(copie)`, `_copy`), returns the canonical
    stem (without suffix, without extension). Returns None if the
    filename is the canonical version.

    Examples:
      "Le_Faire_Causatif (1).docx" -> "Le_Faire_Causatif"
      "cours_articles_-_Copie.docx" -> "cours_articles"
      "homework.docx" -> None
    """
    m = _DUPLICATE_SUFFIX_PATTERN.match(filename)
    if m:
        return m.group("stem")
    return None


def collapse_duplicate_filenames(filenames: list[str]) -> list[str]:
    """Given a list of filenames, return only the canonical (non-
    duplicate-suffix) members. When two filenames share the same
    canonical stem AND both are duplicate-suffixed, keep the earliest
    in sort order (deterministic).

    Examples:
      ["A.docx", "A (1).docx"]              -> ["A.docx"]
      ["A (1).docx", "A (2).docx"]          -> ["A (1).docx"]
      ["A.docx", "A (1).docx", "B.docx"]    -> ["A.docx", "B.docx"]
    """
    canonical_stems: dict[str, list[str]] = {}
    canonical_files: list[str] = []
    for fn in filenames:
        stem = detect_duplicate_suffix(fn)
        if stem is None:
            # Canonical file. Drop any earlier-recorded duplicates
            # with the same stem; keep this one.
            base_stem = fn[:-5] if fn.endswith(".docx") else fn
            canonical_stems.setdefault(base_stem, [])
            canonical_files.append(fn)
        else:
            canonical_stems.setdefault(stem, []).append(fn)

    out: list[str] = []
    seen_stems: set[str] = set()
    for fn in canonical_files:
        base_stem = fn[:-5] if fn.endswith(".docx") else fn
        out.append(fn)
        seen_stems.add(base_stem)
    # Add duplicate-only stems (no canonical original present), taking
    # the alphabetically-first variant.
    for stem, variants in canonical_stems.items():
        if stem in seen_stems:
            continue
        if variants:
            out.append(sorted(variants)[0])
    return sorted(out)


# ── .docx text extraction ────────────────────────────────────────


def extract_docx_text(path: Path) -> str:
    """Read all paragraphs AND table cells from a .docx file as a
    single concatenated string.

    Tables are included because publisher-attribution + ISBN markers
    sometimes live in title-page tables, and the classifier prompt
    benefits from seeing tabular content (it's a signal that the
    table extractor will fire in Phase C).

    Returns an empty string on read failure (corrupted file, lock-
    file, etc.) so the calling pipeline can record + skip without
    crashing.
    """
    try:
        doc = Document(str(path))
    except Exception:
        return ""

    parts: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if row_cells:
                parts.append(" | ".join(row_cells))
    return "\n".join(parts)


# ── Haiku classifier ─────────────────────────────────────────────


CLASSIFIER_SYSTEM_PROMPT = """You are a structured-data extraction classifier for a French-language tutoring archive. The user is Chadi Bakhay, a French tutor with 7,000+ hours teaching anglophones (primarily TCF-Canada candidates). His archive contains ~200 .docx files: tutoring session notes, grammar lessons, exercises, vocabulary lists, and personal pedagogical content.

Your job: read one file's filename and first ~3000 characters of text, and return ONE valid JSON object describing whether the file is a high-yield source of extractable FR/EN lexical chunks for the Le Vocabulaire product.

OUTPUT SCHEMA (strict — return ONLY the JSON, no prose):

{
  "verdict": "HIGH" | "PARTIAL" | "NEEDS_MANUAL" | "SKIP",
  "source_type": "chadi_authored" | "book_lab" | "third_party_publisher_DO_NOT_EXTRACT",
  "content_type": "tutoring_session" | "grammar_lesson" | "vocab_list" | "exercise_set" | "programme" | "publisher_import" | "activity_script" | "other",
  "has_tables": true | false,
  "has_paragraph_glosses": true | false,
  "suggested_topic_slug": "faux_amis" | "calques_anglais" | "prepositions_a_de_dans_par_pour" | "other" | null,
  "confidence": 0.0..1.0,
  "reasoning": "one sentence, max 30 words"
}

VERDICT RUBRIC:
- HIGH: file has clearly structured FR/EN chunk pairs (2-column tables, consistent paragraph glosses like 'faire la grasse matinée (to sleep in)'). Table extractor or regex extractor will yield ≥10 clean chunks.
- PARTIAL: file mixes structured and unstructured content. Some chunks extractable; some need Haiku-assisted extraction.
- NEEDS_MANUAL: file is mostly unstructured prose (a tutoring session transcript, a programme outline). Possible to extract chunks but only with Haiku-assisted prompting.
- SKIP: file is not vocabulary content. Programmes (Avery_Programme_*, Egor_Calendar), activity scripts (one-off lesson plans without paired chunks), copyright violations, corrupted files.

SOURCE_TYPE RUBRIC:
- chadi_authored: tutoring session notes, custom grammar lessons, exercises Chadi wrote for his students. Names like Jack_*, Egor_*, Yarden_*, Avery_*, Andre_*, Matt_*, sumit_*, jack_*. Also generic Chadi-authored grammar lessons (Le_Faire_Causatif, AUTANT, French_Articles_*, cours_articles_*, Subjonctif_*_Egor).
- book_lab: Les_Moules_Complete_Framework.docx (Chadi's Book-Lab pre-publication content). Only this file.
- third_party_publisher_DO_NOT_EXTRACT: published copyrighted FR-as-second-language material. Pre-filter catches CLE/Hachette/Didier/Cambridge by filename + content; this category is here for completeness if the classifier sees a missed attribution.

SUGGESTED_TOPIC_SLUG RUBRIC:
- faux_amis: file focuses on false cognates between FR and EN (faux-amis like 'attendre' (≠ attend), 'librairie' (≠ library)).
- calques_anglais: file focuses on direct English calques in French (anglicismes, structures calquées sur l'anglais).
- prepositions_a_de_dans_par_pour: file focuses on French prepositions (à, de, dans, par, pour, en) — usage rules, exercises, error patterns.
- other: file is about a different topic (subjonctif, conditionnel, articles, cod/coi, faire causatif, conjugation, etc.). Phase 1 only seeds the 3 topics above; chunks suggested as 'other' are kept in the extraction CSV but flagged for Phase 2 deferral.
- null: file is SKIP or has no extractable chunks regardless of topic.

Rules:
1. Return ONLY the JSON object. No prose. No markdown fences.
2. If unsure between two verdicts, pick the more conservative one (PARTIAL over HIGH; NEEDS_MANUAL over PARTIAL; SKIP over NEEDS_MANUAL).
3. confidence is your subjective confidence in the verdict; 0.5 is "could go either way", 0.9 is "very sure".
4. reasoning must be 30 words max and explain the verdict, not summarize the file.
"""


def build_classifier_user_prompt(filename: str, text: str) -> str:
    """Construct the per-file user prompt. Text is truncated at 3000
    chars to keep input tokens predictable."""
    truncated = text[:3000]
    if len(text) > 3000:
        truncated += f"\n\n[... truncated; full length was {len(text)} chars]"
    return (
        f"Filename: {filename}\n\n"
        f"Text (first 3000 chars):\n```\n{truncated}\n```\n\n"
        "Return the JSON verdict object now."
    )
